"""輸出比對報告（xlsx 與 csv）。"""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Any, Iterable

from ..model import CATEGORY_LABEL_ZH, Category, Issue
from .classify import CompareResult
from . import normalize as nz

HEADERS = [
    "分類",
    "對話ID",
    "文本順序",
    "畫面順序",
    "發話者(文本)",
    "發話者(畫面)",
    "譯文(正確答案)",
    "畫面文字(OCR)",
    "相似度",
    "說明",
    "試算表列號",
    "截圖",
]

# 報告底色，純粹為了人眼掃描方便
_CATEGORY_FILL = {
    Category.UNTRANSLATED: "FFC7CE",
    Category.TRUNCATED: "FFD9A0",
    Category.MISMATCH: "FFF2A8",
    Category.ORDER: "D6E4FF",
    Category.MISSING: "E4D6FF",
    Category.SPEAKER: "D6F5E3",
    Category.EXTRA: "E8E8E8",
}

CATEGORY_ORDER = [
    Category.UNTRANSLATED,
    Category.TRUNCATED,
    Category.MISMATCH,
    Category.SPEAKER,
    Category.ORDER,
    Category.MISSING,
    Category.EXTRA,
    Category.PASS,
]


def _row(issue: Issue) -> list[Any]:
    exp, cap = issue.expected, issue.captured
    return [
        CATEGORY_LABEL_ZH[issue.category],
        exp.dialogue_id if exp else "",
        issue.expected_order or "",
        issue.actual_order or "",
        (exp.speaker_en or exp.speaker_zh) if exp else "",
        nz.strip_speaker_id(cap.speaker_text) if cap else "",
        nz.display_key(exp.target_en) if exp else "",
        nz.display_key(cap.body_text) if cap else "",
        round(issue.similarity, 4),
        issue.detail,
        exp.sheet_row if exp else "",
        cap.screenshot if cap else "",
    ]


def _sorted_issues(issues: Iterable[Issue], include_pass: bool) -> list[Issue]:
    rank = {c: i for i, c in enumerate(CATEGORY_ORDER)}
    items = [i for i in issues if include_pass or i.category is not Category.PASS]
    return sorted(
        items,
        key=lambda i: (rank.get(i.category, 99), i.expected_order or i.actual_order or 0),
    )


def write_csv(result: CompareResult, path: str | Path, include_pass: bool = False) -> Path:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    with p.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(HEADERS)
        for issue in _sorted_issues(result.issues, include_pass):
            writer.writerow(_row(issue))
    return p


def write_xlsx(result: CompareResult, path: str | Path, include_pass: bool = True) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()

    # --- 摘要 ---
    ws_sum = wb.active
    ws_sum.title = "摘要"
    counts = result.summary()
    ws_sum.append(["項目", "數量"])
    ws_sum.append(["翻譯文本總句數", len(result.expected)])
    ws_sum.append(["實際錄到句數", len(result.captured)])
    ws_sum.append([])
    ws_sum.append(["分類", "數量"])
    for cat in CATEGORY_ORDER:
        ws_sum.append([CATEGORY_LABEL_ZH[cat], counts.get(cat.value, 0)])
    for cell in ("A1", "B1", "A5", "B5"):
        ws_sum[cell].font = Font(bold=True)
    ws_sum.column_dimensions["A"].width = 22
    ws_sum.column_dimensions["B"].width = 10

    # --- 明細 ---
    ws = wb.create_sheet("明細")
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDDDDD")

    for issue in _sorted_issues(result.issues, include_pass):
        ws.append(_row(issue))
        color = _CATEGORY_FILL.get(issue.category)
        if color:
            ws.cell(row=ws.max_row, column=1).fill = PatternFill("solid", fgColor=color)

    widths = [16, 14, 10, 10, 14, 14, 60, 60, 10, 40, 12, 30]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"

    wb.save(p)
    return p


def print_summary(result: CompareResult) -> None:
    counts = result.summary()
    print("")
    print("=== LQA 比對摘要 ===")
    print(f"翻譯文本總句數 : {len(result.expected)}")
    print(f"實際錄到句數   : {len(result.captured)}")
    for cat in CATEGORY_ORDER:
        n = counts.get(cat.value, 0)
        if n:
            print(f"  {CATEGORY_LABEL_ZH[cat]:<14}{n}")
    print("")
